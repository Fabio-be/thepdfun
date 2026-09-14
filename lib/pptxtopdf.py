import os
from io import BytesIO
from pptx import Presentation
from pptx.util import Emu
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader


def _emu_to_pt(emu):
    """Convertit des EMU (unité PowerPoint) en points (unité PDF)."""
    return Emu(emu).pt


def _get_font_color(run):
    try:
        color = run.font.color
        if color and color.type is not None and color.rgb is not None:
            rgb = str(color.rgb)
            return tuple(int(rgb[i:i+2], 16) / 255 for i in (0, 2, 4))
    except Exception:
        pass
    return (0, 0, 0)  # noir par défaut


def _draw_text_frame(c, shape, slide_height_pt):
    tf = shape.text_frame
    left = _emu_to_pt(shape.left)
    top = _emu_to_pt(shape.top)
    width = _emu_to_pt(shape.width)

    y_cursor = slide_height_pt - top - 14  # position de départ (haut du shape)

    for para in tf.paragraphs:
        x_cursor = left
        align = para.alignment
        line_text_width = 0

        # calcul largeur totale pour alignement centre/droite
        runs_info = []
        for run in para.runs:
            if not run.text:
                continue
            size = run.font.size.pt if run.font.size else 18
            font_name = "Helvetica-Bold" if run.font.bold else "Helvetica"
            if run.font.italic:
                font_name = font_name.replace("Helvetica", "Helvetica-Oblique") if not run.font.bold else "Helvetica-BoldOblique"
            c.setFont(font_name, size)
            text_width = c.stringWidth(run.text, font_name, size)
            runs_info.append((run, font_name, size, text_width))
            line_text_width += text_width

        if align == 2:  # centre
            x_cursor = left + (width - line_text_width) / 2
        elif align == 3:  # droite
            x_cursor = left + width - line_text_width

        max_size = max([r[2] for r in runs_info], default=12)

        for run, font_name, size, text_width in runs_info:
            r, g, b = _get_font_color(run)
            c.setFillColorRGB(r, g, b)
            c.setFont(font_name, size)
            c.drawString(x_cursor, y_cursor, run.text)
            x_cursor += text_width

        y_cursor -= max_size * 1.3  # interligne


def _draw_picture(c, shape, slide_height_pt):
    try:
        image_stream = BytesIO(shape.image.blob)
        img = ImageReader(image_stream)
        left = _emu_to_pt(shape.left)
        top = _emu_to_pt(shape.top)
        width = _emu_to_pt(shape.width)
        height = _emu_to_pt(shape.height)
        y = slide_height_pt - top - height
        c.drawImage(img, left, y, width=width, height=height,
                     preserveAspectRatio=True, mask='auto')
    except Exception:
        pass


def _draw_shape_fill(c, shape, slide_height_pt):
    """Dessine le fond coloré d'une forme (rectangle, etc.) si présent."""
    try:
        fill = shape.fill
        if fill.type is not None and fill.fore_color.rgb is not None:
            rgb = str(fill.fore_color.rgb)
            r, g, b = (int(rgb[i:i+2], 16) / 255 for i in (0, 2, 4))
            left = _emu_to_pt(shape.left)
            top = _emu_to_pt(shape.top)
            width = _emu_to_pt(shape.width)
            height = _emu_to_pt(shape.height)
            y = slide_height_pt - top - height
            c.setFillColorRGB(r, g, b)
            c.rect(left, y, width, height, fill=1, stroke=0)
    except Exception:
        pass


def pptx_to_pdf(input_path, output_path):
    """
    Convertit un fichier PowerPoint (.pptx) en PDF en reconstruisant
    chaque slide (texte, couleurs, images, formes) avec reportlab.
    Positionnement absolu conservé (contrairement au docx).

    Installation: pip install python-pptx reportlab
    """
    input_path = os.path.abspath(input_path)
    output_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    prs = Presentation(input_path)
    slide_width_pt = _emu_to_pt(prs.slide_width)
    slide_height_pt = _emu_to_pt(prs.slide_height)

    c = canvas.Canvas(output_path, pagesize=(slide_width_pt, slide_height_pt))

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                _draw_picture(c, shape, slide_height_pt)
            elif shape.has_text_frame and shape.text_frame.text.strip():
                _draw_shape_fill(c, shape, slide_height_pt)
                _draw_text_frame(c, shape, slide_height_pt)
            else:
                _draw_shape_fill(c, shape, slide_height_pt)

        c.showPage()

    c.save()